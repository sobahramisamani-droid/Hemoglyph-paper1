import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
FIG={"Fig. 1.":"fig/Fig1.png","Fig. 2.":"fig/Fig2_curation.png","Fig. 3.":["fig/Fig3a.png","fig/Fig3b.png"],"Fig. 4.":"fig/Fig3.png","Fig. 5.":"fig/Fig4.png","Fig. 6.":"fig/Fig5.png","Fig. 7.":"fig/Fig6.png"}
doc=Document()
for s in doc.sections: s.left_margin=s.right_margin=Cm(2.5); s.top_margin=s.bottom_margin=Cm(2.5)
n=doc.styles["Normal"]; n.font.name="Times New Roman"; n.font.size=Pt(12); n.paragraph_format.line_spacing=1.5; n.paragraph_format.space_after=Pt(6)
for lvl in (1,2):
    st=doc.styles[f"Heading {lvl}"]; st.font.name="Times New Roman"; st.font.color.rgb=RGBColor(0,0,0); st.font.size=Pt(14 if lvl==1 else 12.5)
def inline(p,text,size=None):
    for pt in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^[^^]+\^)",text):
        if not pt: continue
        if pt.startswith("**"): r=p.add_run(pt[2:-2]); r.bold=True
        elif pt.startswith("*"): r=p.add_run(pt[1:-1]); r.italic=True
        elif pt.startswith("`"): r=p.add_run(pt[1:-1]); r.font.name="Consolas"
        elif pt.startswith("^"): r=p.add_run(pt[1:-1]); r.font.superscript=True
        else: r=p.add_run(pt)
        if size: r.font.size=Pt(size)
def shade(c,f):
    t=c._tc.get_or_add_tcPr(); s=OxmlElement("w:shd"); s.set(qn("w:val"),"clear"); s.set(qn("w:fill"),f); t.append(s)
lines=open("Paper1_Hemoglyph_SampleStyle_v7.md").read().split("\n"); i=0; inrefs=False
while i<len(lines):
    l=lines[i]
    if l.startswith("# "):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; inline(p,l[2:],15); [setattr(r,"bold",True) for r in p.runs]; p.paragraph_format.line_spacing=1.2
    elif l.startswith("## "):
        inrefs = l[3:].strip()=="References"; doc.add_heading(l[3:],1)
    elif l.startswith("### "): doc.add_heading(l[4:],2)
    elif l.startswith("|"):
        rows=[]
        while i<len(lines) and lines[i].startswith("|"): rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i+=1
        hdr,body=rows[0],rows[2:]; t=doc.add_table(rows=1+len(body),cols=len(hdr)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for r,row in enumerate([hdr]+body):
            for c,cell in enumerate((row+[""]*len(hdr))[:len(hdr)]):
                cl=t.cell(r,c); cl.text=""; p=cl.paragraphs[0]; p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(0); inline(p,cell,8.5)
                if r==0: shade(cl,"E7E6E6"); [setattr(x,"bold",True) for x in p.runs]
        doc.add_paragraph(); continue
    elif l.strip()=="": pass
    elif re.match(r"^Fig\. \d+\.",l):
        key=re.match(r"^(Fig\. \d+\.)",l).group(1); fs=FIG.get(key)
        if fs:
            for f in (fs if isinstance(fs,list) else [fs]):
                pp=doc.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pp.add_run().add_picture(f,width=Cm(15.5))
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; inline(p,l,10); [setattr(r,"italic",True) for r in p.runs]
    elif re.match(r"^Table \d+\.",l):
        p=doc.add_paragraph(); inline(p,l,10.5); [setattr(r,"bold",True) for r in p.runs]; p.paragraph_format.keep_with_next=True
    elif inrefs and re.match(r"^\d+\. ",l):
        p=doc.add_paragraph(); inline(p,l,10); p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(2)
    else:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; inline(p,l)
    i+=1
doc.save("Paper1_Hemoglyph_SampleStyle_v7.docx"); print("ok")
