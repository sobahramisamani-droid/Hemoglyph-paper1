import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc=Document()
for s in doc.sections: s.left_margin=s.right_margin=Cm(2.5); s.top_margin=s.bottom_margin=Cm(2.5)
n=doc.styles["Normal"]; n.font.name="Calibri"; n.font.size=Pt(11); n.paragraph_format.line_spacing=1.5; n.paragraph_format.space_after=Pt(6)
def inline(p,text,size=None):
    for pt in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^[^^]+\^)",text):
        if not pt: continue
        if pt.startswith("**"): r=p.add_run(pt[2:-2]); r.bold=True
        elif pt.startswith("*"): r=p.add_run(pt[1:-1]); r.italic=True
        elif pt.startswith("`"): r=p.add_run(pt[1:-1]); r.font.name="Consolas"
        elif pt.startswith("^"): r=p.add_run(pt[1:-1]); r.font.superscript=True
        else: r=p.add_run(pt)
        if size: r.font.size=Pt(size)
def shade(c,f):
    t=c._tc.get_or_add_tcPr(); s=OxmlElement("w:shd"); s.set(qn("w:val"),"clear"); s.set(qn("w:fill"),f); t.append(s)
lines=open("Paper1_Hemoglyph_JBI_EN_v7_R3.md").read().split("\n"); i=0
while i<len(lines):
    l=lines[i]
    if l.startswith("# "):
        p=doc.add_paragraph(); inline(p,l[2:],16); p.runs[0].bold=True; p.paragraph_format.line_spacing=1.2
    elif l.startswith("## "):
        h=doc.add_heading(l[3:],1); h.runs[0].font.color.rgb=RGBColor(0,0,0)
    elif l.startswith("### "):
        h=doc.add_heading(l[4:],2); h.runs[0].font.color.rgb=RGBColor(0x30,0x30,0x30)
    elif l.startswith("|"):
        rows=[]
        while i<len(lines) and lines[i].startswith("|"): rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i+=1
        hdr,body=rows[0],rows[2:]; t=doc.add_table(rows=1+len(body),cols=len(hdr)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for r,row in enumerate([hdr]+body):
            for c,cell in enumerate((row+[""]*len(hdr))[:len(hdr)]):
                cl=t.cell(r,c); cl.text=""; p=cl.paragraphs[0]; p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(0); inline(p,cell,8.5)
                if r==0: shade(cl,"E7E6E6"); [setattr(x,"bold",True) for x in p.runs]
        from docx.oxml.ns import qn as _qn
        from docx.oxml import OxmlElement as _OE
        for ri,tr in enumerate(t.rows):
            trPr=tr._tr.get_or_add_trPr(); cs=_OE("w:cantSplit"); trPr.append(cs)
            if ri==0: th=_OE("w:tblHeader"); trPr.append(th)
            if ri<len(t.rows)-1:
                for cl in tr.cells:
                    for pp in cl.paragraphs: pp.paragraph_format.keep_with_next=True
        doc.add_paragraph(); continue
    elif l.strip()=="" : pass
    elif re.match(r"^\d+\. ",l): p=doc.add_paragraph(); inline(p,l,9.5); p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(2)
    elif l.startswith("**Table "):
        p=doc.add_paragraph(); inline(p,l,10); p.paragraph_format.keep_with_next=True
    elif l.startswith("**Fig. ") and "|" in l and i>0:
        p=doc.add_paragraph(); inline(p,l,10)
    else:
        p=doc.add_paragraph(); inline(p,l)
    i+=1
doc.add_page_break(); h=doc.add_heading("Figures",1); h.runs[0].font.color.rgb=RGBColor(0,0,0)
NPJ=["fig/Fig1.png","fig/Fig3.png","fig/Fig4.png","fig/Fig5.png","fig/Fig6.png","fig/Fig7_llm_v7.png","fig/Fig8_calibration_dca.png"]
LEG={}
for l in lines:
    m=re.match(r"\*\*Fig\. (\d+) \| ([^*]+)\*\*(.*)",l)
    if m: LEG[int(m.group(1))]=(m.group(2).strip(), m.group(3).strip())
for k,f in enumerate(NPJ,1):
    doc.add_paragraph().add_run().add_picture(f,width=Cm(16))
    p=doc.add_paragraph(); r=p.add_run(f"Fig. {k} | {LEG.get(k,('',''))[0]} "); r.bold=True; r.font.size=Pt(9.5)
    inline(p,LEG.get(k,('',''))[1],9.5); p.paragraph_format.line_spacing=1.15
    doc.add_page_break()
doc.save("Paper1_Hemoglyph_JBI_EN_v7_R3.docx"); print("docx ok")
