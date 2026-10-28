"""Shared Elsevier-style formatting helpers for the CMPB manuscript and supplement (python-docx)."""
import re
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_FONT = "Times New Roman"
MONO_FONT = "Courier New"


def _set_font(style_or_run, name, size=None, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    f = style_or_run.font
    f.name = name
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else style_or_run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color


def setup_document(doc, body_size=12, line_spacing=2.0, margins_cm=2.5, line_numbers=True, page_numbers=True,
                   heading_sizes=(14, 12, 12, 12), space_after=6):
    """A4, uniform margins, Times New Roman body, black bold headings, continuous line numbers, centred page numbers."""
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(margins_cm)
        if line_numbers:
            sectPr = s._sectPr
            ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "360")
            sectPr.append(ln)
        if page_numbers:
            p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(); _set_font(r, BODY_FONT, 10)
            for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
                if tag:
                    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
                else:
                    e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text
                r._r.append(e)
    n = doc.styles["Normal"]; _set_font(n, BODY_FONT, body_size)
    n.paragraph_format.line_spacing = line_spacing; n.paragraph_format.space_after = Pt(space_after)
    n.paragraph_format.widow_control = True
    for lvl, size in enumerate(heading_sizes, 1):
        h = doc.styles[f"Heading {lvl}"]
        _set_font(h, BODY_FONT, size, bold=True, italic=(lvl == 3))
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8); h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True; h.paragraph_format.line_spacing = 1.15
    for name in ("List Bullet", "List Number"):
        try: _set_font(doc.styles[name], BODY_FONT, body_size)
        except KeyError: pass
    try:
        _set_font(doc.styles["Title"], BODY_FONT, 16, bold=True)
    except KeyError: pass


def inline(p, text, size=None, mono=MONO_FONT):
    """Render **bold**, *italic*, `code`, ^sup^ markdown fragments into runs."""
    for pt in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^[^^]+\^)", text):
        if not pt: continue
        if pt.startswith("**"): r = p.add_run(pt[2:-2]); r.bold = True
        elif pt.startswith("*"): r = p.add_run(pt[1:-1]); r.italic = True
        elif pt.startswith("`"): r = p.add_run(pt[1:-1]); _set_font(r, mono, (size or 12) - 1)
        elif pt.startswith("^"): r = p.add_run(pt[1:-1]); r.font.superscript = True
        else: r = p.add_run(pt)
        if size: r.font.size = Pt(size)


def add_table(doc, rows, font_size=9, header_fill="F2F2F2"):
    """Editable table, Times, single-spaced, repeat header row, rows not split across pages."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    body = [r for r in rows[1:] if not all(re.fullmatch(r":?-+:?", c) for c in r if c)]
    hdr = rows[0]
    t = doc.add_table(rows=1 + len(body), cols=len(hdr)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate([hdr] + body):
        for c, cell in enumerate((row + [""] * len(hdr))[:len(hdr)]):
            cl = t.cell(r, c); cl.text = ""; p = cl.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            inline(p, cell, font_size)
            if r == 0:
                for x in p.runs: x.bold = True
                tcPr = cl._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), header_fill); tcPr.append(sh)
    for ri, tr in enumerate(t.rows):
        trPr = tr._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0: trPr.append(OxmlElement("w:tblHeader"))
    return t


def heading(doc, text, level):
    h = doc.add_heading(text, level)
    for r in h.runs:
        _set_font(r, BODY_FONT, None, bold=True, italic=(level == 3))
    return h
