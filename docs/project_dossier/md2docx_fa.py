#!/usr/bin/env python3
"""Markdown (Persian, RTL) -> DOCX with Vazirmatn, full RTL/bidi, Persian digits, cover page, TOC, page numbers."""
import re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Vazirmatn"; CODE = "Consolas"; NAVY = RGBColor(0x1F, 0x3A, 0x5F)
FA = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")

def fa_digits(t):
    def rep(m): return m.group(0).translate(FA)
    t = re.sub(r"(?<![A-Za-z0-9_/\.\-@#])\d+(?:[\.,]\d+)*(?![A-Za-z_/\-])", lambda m: rep(m), t)
    t = re.sub(r"(?<=[۰-۹])\.(?=[۰-۹])", "٫", t)
    t = re.sub(r"(?<=[۰-۹]),(?=[۰-۹]{3})", "٬", t)
    return t

doc = Document()
for s in doc.sections:
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.0)

def set_font(rpr_owner, name=FONT):
    rpr = rpr_owner
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), name)

normal = doc.styles["Normal"]; normal.font.name = FONT; normal.font.size = Pt(11)
set_font(normal.element.get_or_add_rPr())
pf = normal.paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.35
for lvl in range(1, 4):
    st = doc.styles[f"Heading {lvl}"]; st.font.name = FONT; st.font.color.rgb = NAVY; set_font(st.element.get_or_add_rPr())
    st.font.size = Pt({1: 17, 2: 13.5, 3: 12}[lvl]); st.paragraph_format.space_before = Pt(14); st.paragraph_format.space_after = Pt(6)

def rtl_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment = align
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); ppr.append(b)

def style_run(r, bold=False, size=None, code=False, italic=False, color=None):
    rpr = r._r.get_or_add_rPr()
    set_font(rpr, CODE if code else FONT)
    r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if code: r.font.size = Pt(max((size or 11) - 2, 8)); r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    if color: r.font.color.rgb = color
    if not code:
        rt = OxmlElement("w:rtl"); rt.set(qn("w:val"), "1"); rpr.append(rt)

def add_inline(p, text, size=None, bold_all=False):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for pt in parts:
        if not pt: continue
        if pt.startswith("**"): style_run(p.add_run(fa_digits(pt[2:-2])), bold=True, size=size)
        elif pt.startswith("`"): style_run(p.add_run(pt[1:-1]), code=True, size=size)
        else: style_run(p.add_run(fa_digits(pt)), bold=bold_all, size=size)

def para(text, size=None, bold=False, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    rtl_par(p, align)
    if indent: p.paragraph_format.right_indent = Cm(indent)
    add_inline(p, text, size, bold_all=bold)
    return p

def heading(text, level):
    p = doc.add_heading(level=level); rtl_par(p, WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run(fa_digits(text)); style_run(r, bold=True, color=NAVY)
    return p

def caption(text):
    p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_after = Pt(12)
    add_inline(p, text, size=9.5, bold_all=True)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), fill); tcpr.append(sh)

def table(rows):
    hdr = rows[0]; body = [r for r in rows[2:] if any(c.strip() for c in r)]
    ncol = len(hdr)
    t = doc.add_table(rows=1 + len(body), cols=ncol); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = t._tbl.tblPr; bv = OxmlElement("w:bidiVisual"); tblpr.append(bv)
    for i, row in enumerate([hdr] + body):
        row = (row + [""] * ncol)[:ncol]
        for j, cell in enumerate(row):
            c = t.cell(i, j); c.text = ""; p = c.paragraphs[0]; rtl_par(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.15
            add_inline(p, cell.strip(), size=9, bold_all=(i == 0))
            if i == 0: shade(c, "D9E2F3")
            elif i % 2 == 0: shade(c, "F5F7FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_field(p, instr):
    r = p.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
    r = p.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate"); r._r.append(f2)
    r = p.add_run("…"); style_run(r)
    r = p.add_run(); f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end"); r._r.append(f3)

lines = open("PROJECT_FULL_FA.md", encoding="utf-8").read().split("\n")
title = lines[0][2:].strip()
for _ in range(6): doc.add_paragraph()
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.CENTER); r = p.add_run(fa_digits(title)); style_run(r, bold=True, size=24, color=NAVY)
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.CENTER); style_run(p.add_run("مرجع یگانهٔ نگارش مقالهٔ نخست — نسخهٔ ۲٫۰"), size=13)
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.CENTER); style_run(p.add_run("موتور v7 · رجیستری ۱۴٫۱٫۰ · محصول ۱٫۳٫۱"), size=11)
for _ in range(10): doc.add_paragraph()
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.CENTER); style_run(p.add_run("سبحان بهرامی"), bold=True, size=13)
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.CENTER); style_run(p.add_run("دبیرستان علامه حلی ۱۰ (سمپاد)، تهران"), size=11)
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.CENTER); style_run(p.add_run("آبان ۱۴۰۵ / اکتبر ۲۰۲۶"), size=11)
doc.add_page_break()
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.RIGHT); style_run(p.add_run("فهرست مطالب (خودکار؛ در Word با F9 به‌روز می‌شود)"), bold=True, size=14, color=NAVY)
p = doc.add_paragraph(); rtl_par(p, WD_ALIGN_PARAGRAPH.RIGHT); add_field(p, 'TOC \\o "1-2" \\h \\z \\u')
doc.add_page_break()

sec = doc.sections[0]
hp = sec.header.paragraphs[0]; rtl_par(hp, WD_ALIGN_PARAGRAPH.RIGHT); style_run(hp.add_run("هموگلیف — سند جامع پروژه"), size=9, color=RGBColor(0x66, 0x66, 0x66))
fp = sec.footer.paragraphs[0]; rtl_par(fp, WD_ALIGN_PARAGRAPH.CENTER); style_run(fp.add_run("صفحهٔ "), size=9); add_field(fp, "PAGE")

i = 1; skip_manual_toc = False
while i < len(lines):
    l = lines[i]
    if l.startswith("## فهرست مطالب"):
        i += 1
        while i < len(lines) and not lines[i].startswith("## "): i += 1
        continue
    if l.startswith("## "):
        doc.add_page_break(); heading(l[3:], 1)
    elif l.startswith("### "): heading(l[4:], 2)
    elif l.startswith("#### "): heading(l[5:], 3)
    elif l.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append([c for c in lines[i].strip().strip("|").split("|")]); i += 1
        table(rows); continue
    elif l.startswith("!["):
        m = re.search(r"\((.*?)\)", l); p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(m.group(1), width=Cm(16.0))
    elif l.strip() == "---": pass
    elif l.startswith("- "): para(l[2:], style="List Bullet")
    elif re.match(r"^\d+\. ", l):
        n, rest = l.split(". ", 1); para(f"{n}. {rest}", indent=0.5)
    elif l.strip() == "": pass
    elif l.startswith("شکل "): caption(l)
    elif l.startswith("پایان سند"): para(l, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    else: para(l)
    i += 1

try:
    lb = doc.styles["List Bullet"]; lb.font.name = FONT; set_font(lb.element.get_or_add_rPr())
except KeyError: pass
doc.save("PROJECT_FULL_FA.docx")
print("ok", os.path.getsize("PROJECT_FULL_FA.docx") // 1024, "KB")
